#!/usr/bin/env python3
"""Generate the portfolio's ATS-readable, two-page professional resume."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path("public/resume/doc");OUT.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.58);sec.bottom_margin=Cm(1.43)
sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
DARK="242D29";GREEN="255D43";MUTED="6F7973"
def rgb(x):return RGBColor.from_string(x)
normal=doc.styles["Normal"];normal.font.name="Carlito";normal.font.size=Pt(9.2);normal.font.color.rgb=rgb(DARK)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1.12
for name,size,color,bold,before,after in [
 ("Name",22,DARK,True,0,4),("ResumeTitle",10.6,GREEN,True,0,4),
 ("SectionTitle",10.1,GREEN,True,11,4),("JobTitle",10,DARK,True,8,1),
 ("JobDate",8.6,MUTED,False,0,3),("BodyText",9.15,DARK,False,0,4),
 ("ContactLine",8.5,"4E6157",False,0,3)]:
 st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal
 st.font.name="Carlito";st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=rgb(color)
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
 st.paragraph_format.keep_with_next=name in ("SectionTitle","JobTitle","JobDate")
def p(text="",style="BodyText"):
 para=doc.add_paragraph(style=style);para.add_run(text);return para
def section(text):
 para=p(text.upper(),"SectionTitle");border=OxmlElement("w:pBdr");bottom=OxmlElement("w:bottom")
 for k,v in (("val","single"),("sz","8"),("color","B3C7B7")):bottom.set(qn("w:"+k),v)
 border.append(bottom);para._p.get_or_add_pPr().append(border)
def label(name,text):
 para=p();r=para.add_run(name);r.bold=True;r.font.color.rgb=rgb(GREEN);para.add_run(text)
def bullet(text):
 para=doc.add_paragraph(style="List Bullet");para.style.font.name="Carlito";para.style.font.size=Pt(9)
 para.paragraph_format.left_indent=Cm(.45);para.paragraph_format.first_line_indent=Cm(-.25)
 para.paragraph_format.space_after=Pt(2.1);para.paragraph_format.line_spacing=1.1
 para.add_run(text)
def job(title,employer,dates,city,items):
 para=doc.add_paragraph(style="JobTitle");para.add_run(title);para.add_run("  |  "+employer).bold=False
 if employer=="MSYS Technologies":para.paragraph_format.page_break_before=True
 p(dates+"  •  "+city,"JobDate")
 for item in items:bullet(item)
p("Amrishkhan Sheik Abdullah","Name")
p("Technology Lead  |  Microservices & Backend Architect  |  Full-Stack Engineering","ResumeTitle")
p("Dubai, UAE  |  +971 52 588 6136  |  amrishkhan05@gmail.com","ContactLine")
p("amrishkhan.dev  |  linkedin.com/in/amrishkhan  |  github.com/amrishkhan05","ContactLine")
section("Professional profile")
p("Technology Lead with 10+ years of experience designing and delivering distributed backend systems, enterprise integrations, secure APIs, payment workflows, and full-stack products. Hands-on architect across Node.js, NestJS, TypeScript, React, and cloud-native services. Experienced in engineering leadership, production operations, and developer productivity tooling.")
section("Selected impact")
for x in [
 "Delivered critical airline passenger journeys spanning check-in, seat selection, ancillaries, and boarding passes, with 95%+ automated test coverage on key modules.",
 "Designed BFF/microservice orchestration, standardized API error handling, and failure-aware integration patterns for enterprise systems.",
 "Integrated Magnati and Network International payments, including webhooks, transaction processing, and reconciliation workflows."]:bullet(x)
section("Technical expertise")
label("Backend & architecture  ","Node.js, NestJS, Express, TypeScript, Java / Spring Boot, REST, GraphQL, WebSockets, microservices, event-driven systems, API orchestration")
label("Frontend  ","React, Next.js, Angular, Vue, Vite, JavaScript, HTML/CSS, Material UI")
label("Data & cloud  ","PostgreSQL, MongoDB, MySQL, SQL Server, AWS (AppSync, Amplify, Lambda), Azure DevOps, Firebase, Linux, Nginx, CI/CD, GitHub Actions")
label("Security & tooling  ","RBAC, authentication, webhook integrations, Jest, Playwright, RAG, MCP, Tauri/Rust, developer tools")
section("Professional experience")
job("Technology Lead Analyst - Backend","techcarrot FZ LLC","Aug 2022 - Present","Dubai, UAE",[
 "Architect and deliver scalable Node.js, NestJS, and TypeScript services supporting airline check-in, seat selection, ancillary services, boarding passes, and downstream enterprise APIs.",
 "Build secure BFF/orchestration layers with validation, API response standardization, retry strategies, fault handling, and asynchronous integration patterns.",
 "Lead payment integration work involving Magnati and Network International, webhook-driven status updates, and reconciliation flows.",
 "Own operational quality through Linux/Nginx administration, monitoring, debugging, release readiness, architecture reviews, code reviews, and mentoring."])
job("Lead Engineer / Full-Stack Developer","MSYS Technologies","May 2021 - Aug 2022","Bangalore, India",[
 "Led full lifecycle technical delivery, from discovery and architecture through implementation, testing, deployment, and production handover.",
 "Built Node.js and TypeScript APIs, React/Next.js experiences, and integrations with Java-based enterprise services.",
 "Promoted reusable engineering patterns, performance improvements, technical reviews, QA collaboration, and delivery standards."])
job("Senior Software Engineer","Full Potential Solutions","Jul 2019 - May 2021","Chennai, India",[
 "Developed GraphQL middleware and cloud-integrated services using AWS AppSync, Amplify, and Lambda to simplify application data access.",
 "Delivered multi-recording audio playlist features, reporting modules, and reusable frontend/application components.",
 "Supported concurrent release cycles, CI/CD delivery, production issue diagnosis, and performance improvements."])
job("Programmer Analyst","Cognizant Technology Solutions","Aug 2016 - Jul 2019","Chennai, India",[
 "Developed enterprise MEAN-stack applications, reusable Angular components, and API-driven business functionality.",
 "Created reusable Angular directives, services, and shared application modules for consistent enterprise interfaces.",
 "Collaborated across requirements, testing, debugging, delivery support, and production maintenance."])
section("Products & open source")
label("Aruvix  |  aruvix.com  ","Built a privacy-first, local-first developer workspace with React, TypeScript, Vite, and Node.js; browser-native JSON utilities, API workflows, format transformations, and developer experience features.")
label("Published npm packages  ","@amrishkhan05/frankly (focused code-change diffs); @amrishkhan05/hallpass (engineering rule enforcement); sql-select-query-generator (configuration-driven SQL SELECT generation).")
section("Education & certifications")
p("B.E., Electrical & Electronics Engineering - St Joseph's College of Engineering, Chennai (2016)")
p("HackerRank: Problem Solving | JavaScript | SQL")
footer=sec.footer.paragraphs[0];footer.text="Amrishkhan Sheik Abdullah  •  amrishkhan.dev"
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.style.font.name="Carlito";footer.style.font.size=Pt(8);footer.style.font.color.rgb=rgb(MUTED)
doc.core_properties.author="Amrishkhan Sheik Abdullah"
doc.core_properties.title="Technology Lead | Microservices & Backend Architect - Resume"
doc.save(OUT/"Amrishkhan-Sheik-Abdullah-Resume.docx")
